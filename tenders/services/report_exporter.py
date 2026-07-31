from io import BytesIO
from html import escape
import re

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from django.utils import timezone
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors


def build_report_pdf(report):
    buffer = BytesIO()
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=_safe_filename(report.tender_project.name),
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='CnTitle', fontName='STSong-Light', fontSize=20, leading=24, spaceAfter=8))
    styles.add(ParagraphStyle(name='CnHeading', fontName='STSong-Light', fontSize=13, leading=16, spaceBefore=4, spaceAfter=2))
    styles.add(ParagraphStyle(name='CnBody', fontName='STSong-Light', fontSize=10, leading=13))

    story = [
        Paragraph('智能投标分析报告', styles['CnTitle']),
        Paragraph(_pdf_rich_text(report.tender_project.name), styles['CnHeading']),
    ]
    story.extend(_pdf_key_value_table(report, styles))
    story.append(Spacer(1, 3))
    _append_pdf_section(story, styles, '报告摘要', [report.summary or '暂无摘要'])
    _append_pdf_section(story, styles, '资质匹配', _format_qualification_match(report))
    _append_pdf_section(story, styles, '业绩匹配', _format_experience_match(report))
    _append_pdf_section(story, styles, '风险清单', _format_risks(report.risks))
    _append_pdf_section(story, styles, '材料清单', _format_material_checklist(report))
    _append_pdf_section(story, styles, '缺失材料', report.missing_materials)
    _append_pdf_section(story, styles, '下一步动作', report.next_actions)
    _append_pdf_section(story, styles, 'Agent执行轨迹', _format_agent_trace(report))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def build_report_docx(report):
    document = Document()
    styles = document.styles
    styles['Normal'].font.name = 'Microsoft YaHei'
    styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    styles['Normal'].font.size = Pt(10.5)

    document.add_heading('智能投标分析报告', level=0)
    document.add_heading(report.tender_project.name, level=1)

    table = document.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for label, value in _report_key_values(report):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value or '-')

    _append_docx_section(document, '报告摘要', [report.summary or '暂无摘要'])
    _append_docx_section(document, '资质匹配', _format_qualification_match(report))
    _append_docx_section(document, '业绩匹配', _format_experience_match(report))
    _append_docx_section(document, '风险清单', _format_risks(report.risks))
    _append_docx_section(document, '材料清单', _format_material_checklist(report))
    _append_docx_section(document, '缺失材料', report.missing_materials)
    _append_docx_section(document, '下一步动作', report.next_actions)
    _append_docx_section(document, 'Agent执行轨迹', _format_agent_trace(report))

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def report_export_filename(report, suffix):
    return f'{_safe_filename(report.tender_project.name)}-AI分析报告.{suffix}'


def _report_key_values(report):
    project = report.tender_project
    raw_report = report.raw_report or {}
    return [
        ('企业', project.company.name if project.company else '-'),
        ('采购方式', project.procurement_method or '-'),
        ('项目类型', project.project_type or '-'),
        ('项目地区', project.region or '-'),
        ('预算金额', _format_currency(project.budget_amount)),
        ('最高限价', _format_currency(raw_report.get('highest_limit_amount'))),
        ('投标截止时间', _format_deadline(project.deadline or raw_report.get('deadline'))),
        ('投标建议', report.get_decision_display()),
        ('匹配评分', report.match_score),
    ]


def _pdf_key_value_table(report, styles):
    rows = [
        [Paragraph(_pdf_rich_text(label), styles['CnBody']), Paragraph(_pdf_rich_text(value or '-'), styles['CnBody'])]
        for label, value in _report_key_values(report)
    ]
    table = Table(rows, colWidths=[35 * mm, 115 * mm])
    table.setStyle(
        TableStyle(
            [
                ('FONTNAME', (0, 0), (-1, -1), 'STSong-Light'),
                ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f0f2f5')),
                ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#d7dce3')),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 8),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]
        )
    )
    return [table]


def _append_pdf_section(story, styles, title, items):
    story.append(Paragraph(_pdf_rich_text(title), styles['CnHeading']))
    values = items or ['暂无']
    for item in values:
        story.append(Paragraph(_pdf_rich_text(f'- {item}'), styles['CnBody']))


def _append_docx_section(document, title, items):
    document.add_heading(title, level=2)
    values = items or ['暂无']
    for item in values:
        document.add_paragraph(str(item), style='List Bullet')


def _format_risks(risks):
    return [
        f"{risk.get('level', '-')}｜{risk.get('type', '-')}：{risk.get('description', '-')}"
        for risk in (risks or [])
    ]


def _format_qualification_match(report):
    match = (report.raw_report or {}).get('qualification_match') or {}
    if not match:
        return []

    items = [
        f"匹配状态：{_status_label(match.get('status'), QUALIFICATION_STATUS_LABELS)}",
        f"资质评分：{match.get('score', '-')}",
    ]
    items.extend(_format_named_list('已匹配资质', match.get('matched')))
    items.extend(_format_named_list('缺失资质', match.get('missing')))
    return items


def _format_experience_match(report):
    match = (report.raw_report or {}).get('experience_match') or {}
    if not match:
        return []

    items = [
        f"业绩评分：{match.get('score', '-')}",
        f"匹配说明：{match.get('summary', '-')}",
    ]
    items.extend(_format_named_list('相似业绩', match.get('matched_cases')))
    return items


def _format_material_checklist(report):
    checklist = (report.raw_report or {}).get('material_checklist') or []
    return [
        f"{item.get('category', '-')}｜{item.get('name', '-')}｜"
        f"{_status_label(item.get('status'), MATERIAL_STATUS_LABELS)}"
        for item in checklist
    ]


def _format_agent_trace(report):
    trace = (report.raw_report or {}).get('agent_trace') or []
    return [
        f"{item.get('agent', '-')}：{_status_label(item.get('status', 'completed'), AGENT_STATUS_LABELS)}"
        for item in trace
    ]


def _format_named_list(label, values):
    if not values:
        return [f"{label}：暂无"]
    return [f"{label}：{'、'.join(str(value) for value in values)}"]


QUALIFICATION_STATUS_LABELS = {
    'matched': '完全匹配',
    'partial': '部分匹配',
    'unknown': '待确认',
}
MATERIAL_STATUS_LABELS = {
    'missing': '缺失',
    'required': '需准备',
    'ready': '已具备',
}
AGENT_STATUS_LABELS = {
    'completed': '已完成',
    'processing': '进行中',
    'failed': '失败',
}


def _status_label(value, mapping):
    key = str(value or '').strip().lower()
    return mapping.get(key, str(value or '待确认'))


def _format_currency(value):
    if value in (None, ''):
        return '-'
    try:
        return f"{float(value):,.0f} 元"
    except (TypeError, ValueError):
        return str(value)


def _format_deadline(value):
    if not value:
        return '-'
    if hasattr(value, 'strftime'):
        if timezone.is_aware(value):
            value = timezone.localtime(value)
        return value.strftime('%Y-%m-%d %H:%M')
    return str(value).replace('T', ' ')[:16]


ASCII_TOKEN_PATTERN = re.compile(r'(?<![A-Za-z0-9])(?:[A-Za-z]+[A-Za-z0-9.-]*|[0-9][0-9,.:%/\-]*)(?![A-Za-z0-9])')


def _pdf_rich_text(value):
    text = escape(str(value))
    return ASCII_TOKEN_PATTERN.sub(lambda match: f'<font name="Helvetica">{match.group(0)}</font>', text)


def _safe_filename(value):
    return ''.join(ch for ch in str(value or 'report') if ch not in r'\/:*?"<>|').strip() or 'report'
